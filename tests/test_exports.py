import unittest
from io import BytesIO
import re

from core import exports


SAMPLE_ROWS = [{
    "question": "What is 20% of 250?",
    "opt_a": "40", "opt_b": "50", "opt_c": "60", "opt_d": "70",
    "answer_index": 1,
    "explanation": "20/100 x 250 = 50.",
    "topic": "Percentage", "difficulty": "easy", "exam": "Banking",
    "review_status": "draft", "source_title": "Percentage class",
    "source_url": "https://www.youtube.com/watch?v=abc12345678",
    "created_at": "2026-08-30T12:00:00",
}, {
    "question": "A is north of B. In which direction is B from A?",
    "opt_a": "North", "opt_b": "South", "opt_c": "East", "opt_d": "West",
    "answer_index": 1,
    "explanation": "B is south of A.",
    "topic": "Direction and Distance", "difficulty": "medium", "exam": "Banking",
    "review_status": "draft", "source_title": "Reasoning class",
    "source_url": "https://www.youtube.com/watch?v=xyz12345678",
    "created_at": "2026-08-30T12:05:00",
}]


class QuestionBankExportTests(unittest.TestCase):
    def test_excel_export_is_readable_and_structured(self):
        from openpyxl import load_workbook

        payload = exports.build_question_bank_xlsx(SAMPLE_ROWS)
        self.assertTrue(payload.startswith(b"PK"))
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook["Faculty Review Bank"]
        self.assertEqual(sheet["A1"].value, "Faculty Review Bank")
        self.assertEqual(sheet["B5"].value, SAMPLE_ROWS[0]["question"])
        self.assertEqual(sheet["G5"].value, "B")
        self.assertEqual(sheet.freeze_panes, "B5")
        self.assertEqual(len(sheet.tables), 1)

    def test_pdf_export_is_readable_and_paginated(self):
        payload = exports.build_question_bank_pdf(SAMPLE_ROWS)
        self.assertTrue(payload.startswith(b"%PDF"))
        self.assertIsNotNone(re.search(rb"/Type\s*/Page\b", payload))
        self.assertGreater(len(payload), 1000)


if __name__ == "__main__":
    unittest.main()
