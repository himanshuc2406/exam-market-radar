"""Polished, cross-platform exports for the temporary Faculty Review Bank."""
from datetime import datetime
from html import escape
from io import BytesIO
import os


EXPORT_COLUMNS = [
    "No", "Question", "Option A", "Option B", "Option C", "Option D",
    "Correct Answer", "Explanation", "Topic", "Difficulty", "Exam",
    "Review Status", "Source Title", "Source URL", "Created At",
]


def _cell(value):
    return "" if value is None else str(value)


def _answer_letter(row):
    try:
        index = int(row.get("answer_index", 0) or 0)
    except (TypeError, ValueError):
        index = 0
    return chr(65 + index) if 0 <= index <= 3 else ""


def _export_row(row, number):
    return [
        number, _cell(row.get("question")), _cell(row.get("opt_a")),
        _cell(row.get("opt_b")), _cell(row.get("opt_c")),
        _cell(row.get("opt_d")), _answer_letter(row),
        _cell(row.get("explanation")), _cell(row.get("topic")),
        _cell(row.get("difficulty")), _cell(row.get("exam")),
        _cell(row.get("review_status") or "draft"),
        _cell(row.get("source_title")), _cell(row.get("source_url")),
        _cell(row.get("created_at")),
    ]


def build_question_bank_xlsx(rows):
    """Return an Excel workbook as bytes, formatted for faculty review."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Faculty Review Bank"
    sheet.sheet_view.showGridLines = False

    violet = "7257D8"
    dark = "16131F"
    white = "FFFFFF"
    pale = "F4F1FF"
    border_color = "DAD5E8"
    green = "DDF5E8"
    yellow = "FFF2CC"
    red = "FDE2E1"

    sheet.merge_cells("A1:O1")
    sheet["A1"] = "Faculty Review Bank"
    sheet["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=white)
    sheet["A1"].fill = PatternFill("solid", fgColor=violet)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34

    sheet.merge_cells("A2:O2")
    sheet["A2"] = (
        f"{len(rows)} draft question(s) | Exported {datetime.now().strftime('%Y-%m-%d %H:%M')} | "
        "Faculty verification required before publishing"
    )
    sheet["A2"].font = Font(name="Aptos", size=10, italic=True, color="5F5870")
    sheet["A2"].fill = PatternFill("solid", fgColor=pale)
    sheet["A2"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 24

    header_row = 4
    for column, label in enumerate(EXPORT_COLUMNS, 1):
        cell = sheet.cell(header_row, column, label)
        cell.font = Font(name="Aptos", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 30

    thin = Side(style="thin", color=border_color)
    for row_number, row in enumerate(rows, header_row + 1):
        values = _export_row(row, row_number - header_row)
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column, value)
            cell.font = Font(name="Aptos", size=9, color="241F2E")
            cell.alignment = Alignment(
                horizontal="center" if column in (1, 7, 10, 12) else "left",
                vertical="top", wrap_text=True,
            )
            cell.border = Border(bottom=thin)
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="FAF9FC")
        sheet.row_dimensions[row_number].height = 54
        source_cell = sheet.cell(row_number, 14)
        if source_cell.value:
            source_cell.hyperlink = source_cell.value
            source_cell.style = "Hyperlink"

    last_row = max(header_row + 1, header_row + len(rows))
    if rows:
        table = Table(displayName="FacultyReviewQuestions", ref=f"A{header_row}:O{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)

    status_validation = DataValidation(
        type="list", formula1='"draft,reviewed,approved,rejected"', allow_blank=False
    )
    sheet.add_data_validation(status_validation)
    status_validation.add(f"L{header_row + 1}:L{last_row}")
    sheet.conditional_formatting.add(
        f"L{header_row + 1}:L{last_row}",
        FormulaRule(formula=[f'$L{header_row + 1}="approved"'], fill=PatternFill("solid", fgColor=green)),
    )
    sheet.conditional_formatting.add(
        f"L{header_row + 1}:L{last_row}",
        FormulaRule(formula=[f'$L{header_row + 1}="draft"'], fill=PatternFill("solid", fgColor=yellow)),
    )
    sheet.conditional_formatting.add(
        f"L{header_row + 1}:L{last_row}",
        FormulaRule(formula=[f'$L{header_row + 1}="rejected"'], fill=PatternFill("solid", fgColor=red)),
    )

    widths = {
        "A": 7, "B": 46, "C": 25, "D": 25, "E": 25, "F": 25,
        "G": 15, "H": 42, "I": 20, "J": 14, "K": 18, "L": 16,
        "M": 28, "N": 38, "O": 21,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "B5"
    sheet.auto_filter.ref = f"A{header_row}:O{last_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"1:{header_row}"
    sheet.sheet_view.zoomScale = 85

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_font_family():
    """Register a Unicode-friendly system font when one is available."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        (r"C:\Windows\Fonts\Nirmala.ttf", r"C:\Windows\Fonts\NirmalaB.ttf"),
        (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
        ("/System/Library/Fonts/Supplemental/Arial Unicode.ttf", None),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ]
    for regular, bold in candidates:
        if not os.path.exists(regular):
            continue
        bold = bold if bold and os.path.exists(bold) else regular
        try:
            pdfmetrics.registerFont(TTFont("ExportSans", regular))
            pdfmetrics.registerFont(TTFont("ExportSans-Bold", bold))
            pdfmetrics.registerFontFamily(
                "ExportSans", normal="ExportSans", bold="ExportSans-Bold",
                italic="ExportSans", boldItalic="ExportSans-Bold",
            )
            return "ExportSans", "ExportSans-Bold"
        except Exception:
            continue
    return "Helvetica", "Helvetica-Bold"


def build_question_bank_pdf(rows):
    """Return a clean, paginated faculty-review PDF as bytes."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    regular, bold = _pdf_font_family()
    output = BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
        topMargin=18 * mm, bottomMargin=17 * mm,
        title="Faculty Review Bank", author="Exam Market Radar",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ExportTitle", parent=styles["Title"], fontName=bold, fontSize=20,
        leading=24, textColor=colors.HexColor("#34255F"), spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        "ExportSubtitle", parent=styles["BodyText"], fontName=regular,
        fontSize=8.5, leading=12, textColor=colors.HexColor("#696273"),
        spaceAfter=12,
    )
    question_style = ParagraphStyle(
        "ExportQuestion", parent=styles["BodyText"], fontName=bold,
        fontSize=10.5, leading=14, textColor=colors.HexColor("#1E1927"),
        spaceAfter=6,
    )
    option_style = ParagraphStyle(
        "ExportOption", parent=styles["BodyText"], fontName=regular,
        fontSize=8.7, leading=11.5, textColor=colors.HexColor("#2E2935"),
    )
    small_style = ParagraphStyle(
        "ExportSmall", parent=styles["BodyText"], fontName=regular,
        fontSize=7.8, leading=10.5, textColor=colors.HexColor("#696273"),
    )
    meta_style = ParagraphStyle(
        "ExportMeta", parent=small_style, fontName=bold,
        textColor=colors.HexColor("#7257D8"), spaceAfter=4,
    )
    center_style = ParagraphStyle(
        "ExportEmpty", parent=subtitle_style, alignment=TA_CENTER, fontSize=11,
    )

    story = [
        Paragraph("Faculty Review Bank", title_style),
        Paragraph(
            escape(
                f"{len(rows)} draft question(s) | Exported {datetime.now().strftime('%Y-%m-%d %H:%M')} | "
                "Verify every answer and explanation before publishing."
            ),
            subtitle_style,
        ),
    ]
    if not rows:
        story.append(Paragraph("No questions available for export.", center_style))

    for number, row in enumerate(rows, 1):
        if number > 1 and (number - 1) % 10 == 0:
            story.append(PageBreak())
        tags = " | ".join(filter(None, [
            f"Q{number}", _cell(row.get("topic")), _cell(row.get("difficulty")),
            _cell(row.get("exam")), f"Status: {_cell(row.get('review_status') or 'draft')}",
        ]))
        story.append(Paragraph(escape(tags), meta_style))
        story.append(Paragraph(escape(_cell(row.get("question"))).replace("\n", "<br/>"), question_style))

        answer = _answer_letter(row)
        options = []
        for letter, key in zip("ABCD", ("opt_a", "opt_b", "opt_c", "opt_d")):
            prefix = "CORRECT - " if letter == answer else ""
            options.append(Paragraph(
                escape(f"{prefix}{letter}) {_cell(row.get(key))}"), option_style
            ))
        option_table = Table(
            [[options[0], options[1]], [options[2], options[3]]],
            colWidths=[88 * mm, 88 * mm], hAlign="LEFT",
        )
        option_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8F6FC")),
            ("BOX", (0, 0), (-1, -1), 0.45, colors.HexColor("#DED8EB")),
            ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E9E5F0")),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(option_table)
        explanation = _cell(row.get("explanation"))
        if explanation:
            story.append(Spacer(1, 4))
            story.append(Paragraph(
                escape(f"Explanation: {explanation}").replace("\n", "<br/>"), small_style
            ))
        source = _cell(row.get("source_title")) or _cell(row.get("source_url"))
        if source:
            story.append(Paragraph(escape(f"Source: {source}"), small_style))
        story.append(Spacer(1, 11))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(colors.HexColor("#7B7485"))
        canvas.drawString(16 * mm, 9 * mm, "Exam Market Radar | Faculty review required")
        canvas.drawRightString(A4[0] - 16 * mm, 9 * mm, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
